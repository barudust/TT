"""Construye la presentación de MICAI 2026 sobre la plantilla oficial.

Uso (desde la raíz del repositorio):

    python PAPER_MICAI/presentacion_micai/fuente/construir_presentacion.py

Salida: PAPER_MICAI/presentacion_micai/MICAI2026_benchmark_trading_signals.pptx

- Portada y cierre con los diseños de la plantilla oficial de MICAI (la
  portada cumple los tres componentes que pide: título, autores con
  institución y logotipos; no se toca el encabezado ni el pie).
- Diapositivas de contenido sobre el diseño "Blank" de la plantilla (pie
  con los patrocinadores).
- Gráficas nativas de PowerPoint (editables), con los números del artículo.
- Notas del orador con el guion en inglés sencillo y su traducción (guion.py).
"""

from __future__ import annotations

import copy
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
from guion import DIAPOSITIVAS, RESPALDO  # noqa: E402

PLANTILLA = AQUI / "plantilla" / "MICAI2026_Template_Author_upt.potx"
LOGOS = AQUI / "logos"
SALIDA = AQUI.parent / "MICAI2026_benchmark_trading_signals.pptx"

# Colores (hex sin #). Los de los modelos son los mismos de las figuras del paper.
TITULO = "1F3C88"
TINTA = "1F2937"
TINTA2 = "4B5563"
TENUE = "8A8F98"
TARJETA = "F2F4F9"
ACENTO = "D81B60"   # magenta de la plantilla, para el número clave
BLANCO = "FFFFFF"
MODELO = {"LR": "2A78D6", "XGBoost": "EB6834", "LSTM": "1BAF7A", "CNN-LSTM": "EDA100", "CNN": "E87BA4"}
NOMBRE = {"LR": "Logistic regression", "XGBoost": "XGBoost", "LSTM": "LSTM", "CNN-LSTM": "CNN-LSTM", "CNN": "1-D CNN"}
CORTO = {"LR": "LR", "XGBoost": "XGBoost", "LSTM": "LSTM", "CNN-LSTM": "CNN-LSTM", "CNN": "1-D CNN"}
ORDEN = ["LR", "XGBoost", "LSTM", "CNN-LSTM", "CNN"]

# Corbel (la de la plantilla) usa cifras de estilo antiguo; para números se usa Calibri.
TITULOS = "Corbel"
CUERPO = "Calibri"

ANCHO, ALTO = 13.333, 7.5
MX = 0.6           # margen lateral
Y_MAX = 6.45       # el pie de patrocinadores empieza en ~6.7 in


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# ---------------------------------------------------------------------------
# Plantilla
# ---------------------------------------------------------------------------
def abrir_plantilla() -> Presentation:
    """Copia el .potx como .pptx (cambia el tipo de contenido) y lo abre."""
    tmp = Path(tempfile.mkdtemp()) / "base.pptx"
    with zipfile.ZipFile(PLANTILLA) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            zout.writestr(item, data)
    return Presentation(str(tmp))


def quitar_diapositivas(prs: Presentation, conservar: set[int]) -> None:
    lst = prs.slides._sldIdLst
    for i, sld in enumerate(list(lst)):
        if i not in conservar:
            prs.part.drop_rel(sld.rId)
            lst.remove(sld)


def diseno(prs: Presentation, nombre: str):
    for lay in prs.slide_layouts:
        if lay.name.strip() == nombre:
            return lay
    raise KeyError(nombre)


# ---------------------------------------------------------------------------
# Primitivas
# ---------------------------------------------------------------------------
def texto(slide, x, y, w, h, parrafos, size=16, color=TINTA, bold=False, align="l", anchor="t",
          nombre=None, espacio=0, interlineado=None, fuente=CUERPO):
    """parrafos: str | list[str | list[(texto, {opciones})]]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if nombre:
        tb.name = nombre
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(parrafos, str):
        parrafos = [parrafos]
    for i, par in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if espacio:
            p.space_after = Pt(espacio)
        if interlineado:
            p.line_spacing = interlineado
        runs = [(par, {})] if isinstance(par, str) else par
        for t, op in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(op.get("size", size))
            f.bold = op.get("bold", bold)
            f.italic = op.get("italic", False)
            f.color.rgb = rgb(op.get("color", color))
            f.name = op.get("fuente", fuente)
    return tb


def caja(slide, x, y, w, h, fill=TARJETA, line=None, radio=0.12, forma=MSO_SHAPE.ROUNDED_RECTANGLE, nombre=None):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    if nombre:
        s.name = nombre
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radio
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def circulo_num(slide, x, y, d, n, fill=TITULO, size=16):
    c = caja(slide, x, y, d, d, fill=fill, forma=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.name = CUERPO
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = rgb(BLANCO)
    return c


def chip(slide, x, y, w, h, t, fill, color=BLANCO, size=14, bold=True):
    c = caja(slide, x, y, w, h, fill=fill, radio=0.5)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = t
    r.font.name = CUERPO
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(color)
    return c


def linea(slide, x1, y1, x2, y2, color=TENUE, ancho=1.5, guiones=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(ancho)
    if guiones:
        ln.line.dash_style = 4  # MSO_LINE.DASH
    return ln


def titulo(slide, t, sub=None):
    texto(slide, MX, 0.38, ANCHO - 2 * MX, 0.75, [[(t, {"bold": True})]], size=34, color=TITULO, nombre="Title",
          fuente=TITULOS)
    if sub:
        texto(slide, MX, 1.08, ANCHO - 2 * MX, 0.45, sub, size=18, color=TINTA2, nombre="Subtitle")


def notas(slide, d: dict) -> None:
    partes = [f"{d['quien'].upper()} · {d['seg']} s", ""]
    partes += d["en"]
    if d.get("pase"):
        partes += ["", "→ " + d["pase"]]
    partes += ["", "ES:"] + d["es"]
    if d.get("pase_es"):
        partes += ["→ " + d["pase_es"]]
    slide.notes_slide.notes_text_frame.text = "\n".join(partes)


# ---------------------------------------------------------------------------
# Gráficas nativas
# ---------------------------------------------------------------------------
def disposicion_manual(chart, x, y, w, h) -> None:
    """Fija el área de trazado (fracciones del marco) para poder dibujar líneas de referencia exactas."""
    pa = chart._chartSpace.chart.plotArea
    lay = pa.find(qn("c:layout"))
    if lay is None:
        lay = etree.Element(qn("c:layout"))
        pa.insert(0, lay)
    for ch in list(lay):
        lay.remove(ch)
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    for tag, val in [("c:layoutTarget", "inner"), ("c:xMode", "edge"), ("c:yMode", "edge"),
                     ("c:x", x), ("c:y", y), ("c:w", w), ("c:h", h)]:
        etree.SubElement(ml, qn(tag)).set("val", str(val))


def estilo_ejes(chart, vmin, vmax, paso, fmt="0.00", size=14):
    chart.font.size = Pt(size)
    chart.font.name = CUERPO
    chart.font.color.rgb = rgb(TINTA2)
    va = chart.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = vmin, vmax, paso
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("E3E6EC")
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.number_format = fmt
    va.tick_labels.number_format_is_linked = False
    ca = chart.category_axis
    ca.format.line.color.rgb = rgb("B8BDC6")
    ca.has_major_gridlines = False
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW


def etiquetas_valor(serie, fmt="0.000", size=13):
    dl = serie.data_labels
    dl.show_value = True
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(size)
    dl.font.bold = True
    dl.font.name = CUERPO
    dl.font.color.rgb = rgb(TINTA)


def colorear_series(chart, claves):
    for s, k in zip(chart.plots[0].series, claves):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(MODELO[k])
        s.format.line.fill.background()
        s.invert_if_negative = False


# ---------------------------------------------------------------------------
# Diapositivas
# ---------------------------------------------------------------------------
def s01_portada(prs, slide) -> None:
    # La plantilla trae dos marcadores (título y subtítulo); se reubican y se llenan.
    ph = {p.placeholder_format.type: p for p in slide.placeholders}
    tit, sub = list(slide.placeholders)[:2]
    for p, (x, y, w, h) in [(tit, (0.85, 1.55, 8.9, 2.45)), (sub, (0.85, 4.05, 8.6, 1.55))]:
        p.left, p.top, p.width, p.height = Inches(x), Inches(y), Inches(w), Inches(h)
    tf = tit.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].runs[0].text = ("A Benchmark of Five Machine-Learning Architectures "
                                     "for Daily Trading-Signal Classification on US Tech Equities")
    r = tf.paragraphs[0].runs[0]
    r.font.size = Pt(30)
    tit.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT

    stf = sub.text_frame
    autores = ["J. B. Polvo Cuatianquiz¹ · A. Reyes Vera¹ · E. Juárez Carbajal²",
               "D. Reyes Ramos¹ · M. Aragón García¹"]
    inst = ["¹ Escuela Superior de Cómputo (ESCOM), Instituto Politécnico Nacional",
            "² Centro de Investigación en Computación (CIC), Instituto Politécnico Nacional"]
    p0 = stf.paragraphs[0]
    p0.runs[0].text = autores[0]
    p0.runs[0].font.size = Pt(18)
    p0._p.addnext(copy.deepcopy(p0._p))
    stf.paragraphs[1].runs[0].text = autores[1]
    p1 = stf.paragraphs[2]
    p1.runs[0].text = inst[0]
    p1.runs[0].font.size = Pt(14)
    p1.runs[0].font.bold = False
    p2 = copy.deepcopy(p1._p)
    p1._p.addnext(p2)
    stf.paragraphs[3].runs[0].text = inst[1]
    sub.top = Inches(4.0)
    sub.height = Inches(1.75)
    for p in stf.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        # sin viñetas
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        if pPr.find(qn("a:buNone")) is None:
            etree.SubElement(pPr, qn("a:buNone"))
        p.space_after = Pt(4)

    # Logotipos sobre una tarjeta blanca (la plantilla pide logotipos de la institución)
    caja(slide, 9.85, 4.1, 2.75, 1.45, fill=BLANCO, radio=0.18, nombre="Logos")
    slide.shapes.add_picture(str(LOGOS / "ipn.png"), Inches(10.1), Inches(4.22), height=Inches(1.2))
    slide.shapes.add_picture(str(LOGOS / "escom.png"), Inches(11.25), Inches(4.25), height=Inches(1.15))


def s02_pregunta(prs, slide) -> None:
    titulo(slide, "Do deep networks beat simple models?",
           "Daily BUY / HOLD / SELL signals for seven large US tech stocks")
    texto(slide, MX, 1.85, 8, 0.4, "What is often missing in the literature", size=18, color=TINTA, bold=True)
    tarjetas = [
        ("Only 1–2 models", "Few studies compare deep models against a simple linear baseline."),
        ("Accuracy or RMSE", "These metrics do not tell if a trading strategy makes money."),
        ("Untested feature filters", "Statistical pre-selection is applied without checking that it helps."),
    ]
    w, gap = 3.85, 0.3
    for i, (h, d) in enumerate(tarjetas):
        x = MX + i * (w + gap)
        caja(slide, x, 2.45, w, 2.35, fill=TARJETA)
        circulo_num(slide, x + 0.3, 2.7, 0.55, i + 1, fill=ACENTO)
        texto(slide, x + 0.3, 3.45, w - 0.6, 0.45, h, size=20, color=TINTA, bold=True)
        texto(slide, x + 0.3, 3.95, w - 0.6, 0.8, d, size=16, color=TINTA2)
    caja(slide, MX, 5.2, ANCHO - 2 * MX, 1.0, fill="E8F0FC")
    texto(slide, MX + 0.35, 5.2, ANCHO - 2 * MX - 0.7, 1.0,
          [[("Our goal: ", {"bold": True, "color": TITULO}),
            ("a fair five-model benchmark, scored as classification and as a trading strategy.", {})]],
          size=20, anchor="m")


def s03_datos(prs, slide) -> None:
    titulo(slide, "One signal per day, seven tech stocks", "Free daily data · adaptive three-class label")
    # Izquierda: acciones y datos
    texto(slide, MX, 1.85, 5.8, 0.4, "Stocks", size=18, bold=True)
    tick = ["AAPL", "NVDA", "TSLA", "AMZN", "MSFT", "GOOGL", "META"]
    for i, t in enumerate(tick):
        fila, col = divmod(i, 4)
        chip(slide, MX + col * 1.45, 2.35 + fila * 0.62, 1.3, 0.48, t, fill=TARJETA, color=TITULO, size=16)
    datos = [
        [("Yahoo Finance, ", {"bold": True}), ("free, 2013–2025", {})],
        [("61 technical features ", {"bold": True}), ("in 9 families", {})],
        [("Market context: ", {"bold": True}), ("S&P 500 (SPY) and VIX", {})],
    ]
    for i, d in enumerate(datos):
        circulo_num(slide, MX, 3.85 + i * 0.72, 0.42, "✓", fill=MODELO["LR"], size=14)
        texto(slide, MX + 0.6, 3.85 + i * 0.72, 5.3, 0.42, [d], size=18, anchor="m")

    # Derecha: la etiqueta
    x0 = 7.0
    texto(slide, x0, 1.85, 5.7, 0.4, "Label for day t (tomorrow's return vs. last 252 days)", size=18, bold=True)
    zonas = [("BUY", "top 30 %", "1F8A4C", 1.05), ("HOLD", "middle 40 %", "C27C0E", 1.35), ("SELL", "bottom 30 %", "C0392B", 1.05)]
    y = 2.4
    for nom, desc, col, h in zonas:
        caja(slide, x0, y, 1.6, h - 0.08, fill=col, radio=0.12)
        texto(slide, x0, y, 1.6, h - 0.08, [[(nom, {"bold": True, "color": BLANCO})]], size=20, align="c", anchor="m")
        texto(slide, x0 + 1.85, y, 3.8, h - 0.08, desc, size=18, color=TINTA, anchor="m")
        y += h
    texto(slide, x0, 5.95, 5.7, 0.5,
          [[("Adapts to volatility · uses only past data · ≈ 30 / 40 / 30 %", {"color": TINTA2})]], size=16)


def _barra_tiempo(slide, x0, y, anios, ancho_anio, tramos, alto=0.36):
    for a1, a2, col, et, tc in tramos:
        xs = x0 + (a1 - anios[0]) * ancho_anio
        w = (a2 - a1 + 1) * ancho_anio - 0.05
        caja(slide, xs, y, w, alto, fill=col, radio=0.25)
        if et:
            texto(slide, xs, y, w, alto, [[(et, {"bold": True, "color": tc})]], size=13, align="c", anchor="m")


def s04_protocolo(prs, slide) -> None:
    titulo(slide, "Five models, one protocol", "Same data, same features, same test year")
    # Modelos
    texto(slide, MX, 1.8, 4, 0.4, "Classical", size=16, color=TINTA2, bold=True)
    texto(slide, 4.35, 1.8, 6, 0.4, "Deep learning", size=16, color=TINTA2, bold=True)
    for i, k in enumerate(["LR", "XGBoost"]):
        chip(slide, MX + i * 1.85, 2.2, 1.7, 0.52, CORTO[k] if k != "LR" else "Logistic reg.", fill=MODELO[k], size=15)
    for i, k in enumerate(["LSTM", "CNN", "CNN-LSTM"]):
        chip(slide, 4.35 + i * 1.85, 2.2, 1.7, 0.52, CORTO[k], fill=MODELO[k], color=TINTA if k == "CNN-LSTM" else BLANCO, size=15)
    # Particiones
    anios = list(range(2014, 2026))
    x0, aw = 1.75, 0.6
    texto(slide, MX, 3.15, 6, 0.4, "Three splits, same test year", size=16, color=TINTA2, bold=True)
    for i, a in enumerate(anios):
        if a % 2 == 0 or a == 2025:
            texto(slide, x0 + i * aw, 3.55, aw, 0.3, str(a), size=12, color=TENUE, align="c")
    filas = [("A · 10 y", 2014), ("B · 6 y", 2018), ("C · 4 y", 2020)]
    for j, (et, ini) in enumerate(filas):
        y = 3.92 + j * 0.5
        texto(slide, MX, y, 1.1, 0.36, et, size=14, color=TINTA, bold=(et.startswith("B")), anchor="m")
        _barra_tiempo(slide, x0, y, anios, aw, [(ini, 2023, MODELO["LR"], "train", BLANCO),
                                                (2024, 2024, "F1C36B", "val", TINTA),
                                                (2025, 2025, MODELO["XGBoost"], "test", BLANCO)])
    # Derecha: números
    xr = 9.45
    caja(slide, xr, 3.15, 3.28, 3.2, fill=TARJETA)
    texto(slide, xr + 0.3, 3.3, 2.7, 0.8, [[("120", {"bold": True, "color": ACENTO, "size": 44})]], size=44)
    texto(slide, xr + 0.3, 4.1, 2.75, 0.75, ["runs: 5 models × 3 splits", "× (1 global + 7 per stock)"], size=14, color=TINTA2)
    texto(slide, xr + 0.3, 5.0, 2.75, 1.2,
          [[("Main metric: ", {"bold": True}), ("macro-F1", {})],
           [("Plus: ", {"bold": True}), ("trading simulation (Sharpe, profit factor)", {})]],
          size=14, color=TINTA, espacio=4)


def s05_resultado1(prs, slide) -> None:
    titulo(slide, "Result 1 · Logistic regression leads on average", "Macro-F1 on the 2025 test year, global models (Table 3)")
    datos = {"LR": (0.411, 0.385, 0.401), "XGBoost": (0.396, 0.386, 0.373), "LSTM": (0.389, 0.370, 0.370),
             "CNN-LSTM": (0.376, 0.384, 0.371), "CNN": (0.361, 0.343, 0.373)}
    cd = CategoryChartData()
    cd.categories = ["Exp A · 10 years", "Exp B · 6 years", "Exp C · 4 years"]
    for k in ORDEN:
        cd.add_series(CORTO[k], datos[k])
    x, y, w, h = MX, 1.65, 8.4, 4.8
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    vmin, vmax = 0.30, 0.42
    estilo_ejes(ch, vmin, vmax, 0.02)
    colorear_series(ch, ORDEN)
    ch.plots[0].gap_width = 70
    ch.plots[0].overlap = -8
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(14)
    px, py, pw, ph = 0.08, 0.04, 0.82, 0.74
    disposicion_manual(ch, px, py, pw, ph)
    etiquetas_valor(ch.plots[0].series[0])
    # Línea del azar (1/3) en la posición exacta del área de trazado
    yy = y + h * (py + ph * (vmax - 1 / 3) / (vmax - vmin))
    linea(slide, x + w * px, yy, x + w * (px + pw) + 0.1, yy, color="6B7280", ancho=1.5, guiones=True)
    texto(slide, x + w * (px + pw) + 0.15, yy - 0.27, 0.9, 0.55, ["random", "0.333"], size=12, color="6B7280")

    xr = 9.35
    caja(slide, xr, 1.75, 3.38, 2.1, fill=TARJETA)
    texto(slide, xr + 0.3, 1.9, 2.9, 0.85, [[("0.399", {"bold": True, "color": MODELO["LR"], "size": 44})]], size=44)
    texto(slide, xr + 0.3, 2.75, 2.9, 1.0, "average macro-F1 of logistic regression (best)", size=15, color=TINTA2)
    caja(slide, xr, 4.05, 3.38, 2.35, fill=TARJETA)
    texto(slide, xr + 0.3, 4.2, 2.9, 0.4, "Exp B · trading simulation", size=15, color=TINTA, bold=True)
    texto(slide, xr + 0.3, 4.65, 2.9, 1.7,
          [[("Sharpe  ", {"color": TINTA2}), ("+0.76", {"bold": True, "color": MODELO["LR"], "size": 22})],
           [("Profit factor  ", {"color": TINTA2}), ("1.21", {"bold": True, "color": MODELO["LR"], "size": 22})],
           [("both the highest of the five", {"color": TINTA2, "size": 13})]],
          size=16, espacio=6)


def s06_resultado2(prs, slide) -> None:
    titulo(slide, "Result 2 · Do not pre-filter the features",
           "Δ macro-F1 = all 61 features − filtered features (Table 2)")
    delta = {"LR": (0.024, 0.026, 0.021), "XGBoost": (0.009, 0.021, 0.022), "LSTM": (0.039, 0.036, -0.006),
             "CNN-LSTM": (-0.011, 0.020, -0.026), "CNN": (0.035, -0.025, 0.011)}
    cd = CategoryChartData()
    cd.categories = ["Exp A", "Exp B", "Exp C"]
    for k in ORDEN:
        cd.add_series(CORTO[k], delta[k])
    x, y, w, h = MX, 1.65, 8.4, 4.8
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    estilo_ejes(ch, -0.03, 0.045, 0.015, fmt="+0.000;-0.000;0")
    colorear_series(ch, ORDEN)
    ch.plots[0].gap_width = 70
    ch.plots[0].overlap = -8
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(14)
    disposicion_manual(ch, 0.1, 0.04, 0.88, 0.74)

    xr = 9.35
    caja(slide, xr, 1.75, 3.38, 2.45, fill=TARJETA)
    texto(slide, xr + 0.3, 1.9, 2.9, 0.9, [[("11 / 15", {"bold": True, "color": ACENTO, "size": 44})]], size=44)
    texto(slide, xr + 0.3, 2.8, 2.9, 1.3, "cases where keeping all 61 features is better", size=15, color=TINTA2)
    caja(slide, xr, 4.4, 3.38, 2.0, fill=TARJETA)
    texto(slide, xr + 0.3, 4.55, 2.9, 1.8,
          [[("Filters tested: ", {"bold": True}), ("Pearson + VIF, SHAP, Spearman", {})],
           [("The model's own regularization selects better than an external filter.", {"color": TINTA2})]],
          size=15, espacio=6)


def s07_resultado3(prs, slide) -> None:
    titulo(slide, "Result 3 · The stock matters more than the model",
           "Share of the variation of 105 per-stock Sharpe ratios (Section 7.3)")
    # Dos números grandes
    for i, (num, et, col) in enumerate([("15.5 %", "explained by the stock", TITULO),
                                        ("3.4 %", "explained by the model", MODELO["XGBoost"])]):
        x = MX + i * 4.2
        caja(slide, x, 1.85, 3.9, 2.2, fill=TARJETA)
        texto(slide, x + 0.35, 2.0, 3.3, 1.1, [[(num, {"bold": True, "color": col, "size": 54})]], size=54)
        texto(slide, x + 0.35, 3.15, 3.3, 0.6, et, size=18, color=TINTA2)
    caja(slide, 9.0, 1.85, 3.73, 2.2, fill="E8F0FC")
    texto(slide, 9.3, 1.85, 3.2, 2.2, [[("≈ 4.5×", {"bold": True, "color": TITULO, "size": 44})],
                                       [("more from the stock than from the model", {"color": TINTA2, "size": 16})]],
          size=16, anchor="m")
    # Barra al 100 %
    partes = [("Stock 15.5 %", 15.5, "1C5CAB", BLANCO), ("Exp. 7.0 %", 7.0, "86B6EF", TINTA),
              ("Model 3.4 %", 3.4, MODELO["XGBoost"], BLANCO), ("Interaction and noise 74.1 %", 74.1, "D4D7DC", TINTA)]
    x, w_total, y = MX, ANCHO - 2 * MX, 4.6
    texto(slide, MX, 4.25, 6, 0.35, "Decomposition of the total sum of squares", size=14, color=TINTA2)
    for et, v, col, tc in partes:
        w = w_total * v / 100
        caja(slide, x, y, w - 0.03, 0.6, fill=col, forma=MSO_SHAPE.RECTANGLE)
        if v > 10:
            texto(slide, x, y, w, 0.6, [[(et, {"bold": True, "color": tc})]], size=14, align="c", anchor="m")
        x += w
    texto(slide, MX + w_total * 0.155, 5.3, 4, 0.35, "Exp. 7.0 % · Model 3.4 %", size=13, color=TINTA2)
    texto(slide, MX, 5.75, ANCHO - 2 * MX, 0.5,
          "AMZN does well with almost any model; GOOGL does badly with almost any model.", size=16, color=TINTA)


def s08_protocolo_corregido(prs, slide) -> None:
    titulo(slide, "Was it fair? We rebuilt the protocol", "Reviewers' request → Section 8 of the paper")
    anios = list(range(2018, 2026))
    x0, aw = 2.15, 0.82
    for i, a in enumerate(anios):
        texto(slide, x0 + i * aw, 1.85, aw, 0.3, str(a), size=13, color=TENUE, align="c", bold=True)
    filas = [
        ("dev", "all decisions", [(2018, 2022, MODELO["LR"], "train", BLANCO), (2023, 2023, "A9C9F2", "epoch", TINTA),
                                  (2024, 2024, "F1C36B", "config.", TINTA), (2025, 2025, "E5E7EB", "not used", TENUE)]),
        ("final", "frozen config.", [(2018, 2023, MODELO["LR"], "train", BLANCO), (2024, 2024, "A9C9F2", "epoch", TINTA),
                                     (2025, 2025, MODELO["XGBoost"], "test ×1", BLANCO)]),
    ]
    for j, (nom, sub, tramos) in enumerate(filas):
        y = 2.3 + j * 0.95
        texto(slide, MX, y - 0.05, 1.5, 0.35, nom, size=18, bold=True)
        texto(slide, MX, y + 0.3, 1.5, 0.3, sub, size=12, color=TINTA2)
        _barra_tiempo(slide, x0, y, anios, aw, tramos, alto=0.5)
    checks = [("Same 1,736 test rows", "for all five models (before: networks lost the first 20–60 days)"),
              ("Same Optuna search", "150 trials (LR, XGBoost) · 80 trials, 17–20 hyper-parameters (networks)"),
              ("No decision uses 2025", "the test year is evaluated only once")]
    for i, (h, d) in enumerate(checks):
        y = 4.4 + i * 0.68
        circulo_num(slide, MX, y + 0.04, 0.46, "✓", fill="1F8A4C", size=15)
        texto(slide, MX + 0.65, y, 11.4, 0.55, [[(h + "  ", {"bold": True}), (d, {"color": TINTA2, "size": 15})]],
              size=18, anchor="m")


def s09_resultado4(prs, slide) -> None:
    titulo(slide, "Result 4 · The gap does not close",
           "Exp B, corrected protocol: macro-F1 with 95 % block-bootstrap intervals (Table 6)")
    ci = {"LR": (0.404, 0.379, 0.428), "CNN-LSTM": (0.372, 0.346, 0.395), "LSTM": (0.369, 0.341, 0.395),
          "XGBoost": (0.359, 0.334, 0.381), "CNN": (0.359, 0.332, 0.386)}
    orden = ["LR", "CNN-LSTM", "LSTM", "XGBoost", "CNN"]
    # Barra flotante: base invisible hasta el límite inferior; tramo oscuro hasta el F1;
    # tramo claro hasta el límite superior. El cambio de tono marca el F1.
    cd = CategoryChartData()
    cd.categories = [CORTO[k] + chr(10) + f"{ci[k][0]:.3f}" for k in orden]
    cd.add_series("base", [ci[k][1] for k in orden])
    cd.add_series("lower", [ci[k][0] - ci[k][1] for k in orden])
    cd.add_series("upper", [ci[k][2] - ci[k][0] for k in orden])
    x, y, w, h = MX, 1.65, 8.4, 4.8
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    vmin, vmax = 0.32, 0.44
    estilo_ejes(ch, vmin, vmax, 0.02)
    ch.has_legend = False
    ch.plots[0].gap_width = 140
    base, baja, alta = ch.plots[0].series
    base.format.fill.background()
    base.format.line.fill.background()
    for i, k in enumerate(orden):
        for serie, transp in ((baja, 0), (alta, 55)):
            pt = serie.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(MODELO[k])
            if transp:
                alpha = etree.SubElement(pt.format.fill._xPr.find(qn("a:solidFill"))[0], qn("a:alpha"))
                alpha.set("val", str((100 - transp) * 1000))
            pt.format.line.color.rgb = rgb(BLANCO)
            pt.format.line.width = Pt(1.5)
    px, py, pw, ph = 0.1, 0.04, 0.8, 0.74
    disposicion_manual(ch, px, py, pw, ph)
    yy = y + h * (py + ph * (vmax - 1 / 3) / (vmax - vmin))
    linea(slide, x + w * px, yy, x + w * (px + pw) + 0.1, yy, color="6B7280", ancho=1.5, guiones=True)
    texto(slide, x + w * (px + pw) + 0.15, yy - 0.27, 0.9, 0.55, ["random", "0.333"], size=12, color="6B7280")

    xr2 = 9.35
    caja(slide, xr2, 1.75, 3.38, 2.35, fill=TARJETA)
    texto(slide, xr2 + 0.3, 1.9, 2.8, 0.85, [[("0.404", {"bold": True, "color": MODELO["LR"], "size": 44})]], size=44)
    texto(slide, xr2 + 0.3, 2.75, 2.8, 1.3, "LR [0.379, 0.428] · significantly better than each of the other four",
          size=15, color=TINTA2)
    caja(slide, xr2, 4.3, 3.38, 2.1, fill=TARJETA)
    texto(slide, xr2 + 0.3, 4.45, 2.8, 1.9,
          [[("Networks: ", {"bold": True}), ("validation loss never below ln 3 = 1.099", {})],
           [("592 k → 10 k parameters: ", {"bold": True}), ("same result", {})]],
          size=15, color=TINTA, espacio=8)
    texto(slide, MX, 6.15, 8.3, 0.3, "Dark part: lower bound → macro-F1 · light part: macro-F1 → upper bound",
          size=12, color=TINTA2)


def s10_por_que(prs, slide) -> None:
    titulo(slide, "Why does a linear model win here?")
    razones = [("Small data", "≈ 10,500 training rows vs. 10⁵–10⁶ network parameters"),
               ("Noisy label", "one-day direction is close to random"),
               ("Rich features", "indicators already summarize many days")]
    w, gap = 3.85, 0.3
    for i, (h, d) in enumerate(razones):
        x = MX + i * (w + gap)
        caja(slide, x, 1.45, w, 2.3, fill=TARJETA)
        circulo_num(slide, x + 0.3, 1.7, 0.55, i + 1, fill=MODELO["LR"])
        texto(slide, x + 0.3, 2.45, w - 0.6, 0.45, h, size=20, bold=True)
        texto(slide, x + 0.3, 2.95, w - 0.6, 0.75, d, size=16, color=TINTA2)
    texto(slide, MX, 4.05, 8, 0.4, "Limitations", size=18, bold=True)
    lims = ["One test year (2025)", "No transaction costs", "Seven stocks, one sector"]
    for i, t in enumerate(lims):
        chip(slide, MX + i * 4.15, 4.55, 3.85, 0.6, t, fill="FDEEE6", color="9A3412", size=16)
    texto(slide, MX, 5.45, ANCHO - 2 * MX, 0.8,
          [[("Same conclusion in the literature: ", {"bold": True}),
            ("tree and linear models match deep networks on tabular data (Grinsztajn et al., 2022); "
             "networks win with ~30,000 stocks and 60 years (Gu, Kelly and Xiu, 2020).", {"color": TINTA2})]],
          size=15)


def s11_cierre(prs, slide) -> None:
    texto(slide, 0.85, 1.45, 8.5, 0.7, [[("Takeaways", {"bold": True, "color": BLANCO})]], size=36)
    msgs = ["A well-regularized logistic regression is a strong and cheap baseline.",
            "Do not filter the features before training.",
            "The stock matters more than the model."]
    for i, m in enumerate(msgs):
        y = 2.35 + i * 0.85
        circulo_num(slide, 0.85, y, 0.55, i + 1, fill=BLANCO)
        c = slide.shapes[-1]
        c.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(TITULO)
        texto(slide, 1.65, y, 7.9, 0.55, [[(m, {"bold": True, "color": BLANCO})]], size=20, anchor="m")
    texto(slide, 0.85, 5.0, 8.5, 0.6, [[("Thank you · Questions?", {"bold": True, "color": BLANCO})]], size=28)
    texto(slide, 0.85, 5.6, 8.5, 0.4, [[("Code and data: github.com/barudust/TT", {"color": BLANCO})]], size=16)


# --- Respaldo -----------------------------------------------------------------
def tabla(slide, x, y, w, filas, anchos, size=13, encabezado=TITULO, resaltar_fila=None):
    t = slide.shapes.add_table(len(filas), len(filas[0]), Inches(x), Inches(y), Inches(w), Inches(0.38 * len(filas))).table
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.text = str(v)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.runs[0]
            r.font.size = Pt(size)
            r.font.name = CUERPO
            r.font.bold = i == 0 or (resaltar_fila is not None and i == resaltar_fila)
            r.font.color.rgb = rgb(BLANCO if i == 0 else TINTA)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(encabezado if i == 0 else ("EEF3FB" if i == resaltar_fila else ("FFFFFF" if i % 2 else "F6F7FA")))
            c.margin_top = c.margin_bottom = Inches(0.04)
    return t


def b12_tabla6(prs, slide) -> None:
    titulo(slide, "Backup · Corrected protocol, Exp B (Table 6)")
    filas = [["Model", "Trials", "Hyper-param.", "Macro-F1", "95 % CI", "Δ vs. Sec. 7", "Sharpe"],
             ["LR", "150", "5", "0.404", "[0.379, 0.428]", "+0.019", "+0.895"],
             ["CNN-LSTM", "80", "19", "0.372", "[0.346, 0.395]", "−0.011", "−0.074"],
             ["LSTM", "80", "17", "0.369", "[0.341, 0.395]", "−0.001", "−0.345"],
             ["XGBoost", "150", "9", "0.359", "[0.334, 0.381]", "−0.027", "+0.462"],
             ["1-D CNN", "80", "20", "0.359", "[0.332, 0.386]", "+0.016", "−0.459"]]
    tabla(slide, MX, 1.4, 12.1, filas, [2.2, 1.3, 1.6, 1.6, 2.2, 1.6, 1.6], size=15, resaltar_fila=1)
    texto(slide, MX, 4.25, ANCHO - 2 * MX, 1.9,
          [[("Paired differences of LR (all exclude zero): ", {"bold": True}),
            ("+0.045 XGBoost · +0.035 LSTM · +0.031 CNN-LSTM · +0.045 1-D CNN.", {})],
           [("None of the six differences among the other four models is significant.", {"color": TINTA2})],
           [("Biggest lever for the networks: ", {"bold": True}),
            ("choosing the epoch by validation F1 (+0.059 for the LSTM). Retraining them on train + validation makes them worse (−0.024 to −0.047).", {"color": TINTA2})]],
          size=16, espacio=8)


def b13_despues(prs, slide) -> None:
    titulo(slide, "Backup · Since submission (not in the paper)")
    filas = [["Check", "Result"],
             ["Six test years, 2020–2025 (fixed configurations)", "LR 0.360 · XGBoost 0.356 (difference not significant) · networks 0.326–0.338 (significantly lower)"],
             ["Transaction costs, 10 basis points (Exp B, 2025)", "Sharpe: LR +0.89 → +0.62 · XGBoost +0.46 → +0.01 (turnover 0.38 vs. 0.65)"],
             ["Unseen data, Jan–Oct 2026", "LR first in macro-F1 (0.346), not significant vs. XGBoost, LSTM, CNN-LSTM; 1-D CNN and LSTM had a higher Sharpe"],
             ["Eight metrics on the same rows (2025)", "LR best in all eight in A, B and C (F1, accuracy, balanced acc., kappa, AUC, log-loss, Sharpe 0 / 10 bp)"]]
    tabla(slide, MX, 1.4, 12.1, filas, [4.3, 7.8], size=15)
    texto(slide, MX, 4.55, ANCHO - 2 * MX, 0.8,
          "Same conclusion: classical regularized models match or beat the networks; LR vs. XGBoost is a tie over six years.",
          size=16, color=TINTA2)


def b14_tablas(prs, slide) -> None:
    titulo(slide, "Backup · Benchmark tables (Section 7)")
    t3 = [["Global", "A", "B", "C", "Avg"],
          ["LR", "0.411", "0.385", "0.401", "0.399"], ["XGBoost", "0.396", "0.386", "0.373", "0.385"],
          ["LSTM", "0.389", "0.370", "0.370", "0.376"], ["CNN-LSTM", "0.376", "0.384", "0.371", "0.377"],
          ["1-D CNN", "0.361", "0.343", "0.373", "0.359"]]
    t5 = [["Per stock", "A", "B", "C", "Avg"],
          ["LR", "0.411", "0.392", "0.359", "0.387"], ["XGBoost", "0.376", "0.370", "0.352", "0.366"],
          ["CNN-LSTM", "0.368", "0.342", "0.356", "0.355"], ["LSTM", "0.314", "0.353", "0.326", "0.331"],
          ["1-D CNN", "0.378", "0.332", "0.280", "0.330"]]
    tabla(slide, MX, 1.35, 5.9, t3, [1.9, 1.0, 1.0, 1.0, 1.0], size=14, resaltar_fila=1)
    tabla(slide, 6.83, 1.35, 5.9, t5, [1.9, 1.0, 1.0, 1.0, 1.0], size=14, resaltar_fila=1)
    t4 = [["Exp B economics", "F1", "Win rate", "Profit factor", "Max DD", "Sharpe"],
          ["LR", "0.385", "0.497", "1.210", "−0.470", "+0.755"], ["XGBoost", "0.386", "0.500", "1.083", "−0.432", "+0.345"],
          ["CNN-LSTM", "0.384", "0.505", "1.041", "−0.403", "+0.147"], ["LSTM", "0.370", "0.491", "1.120", "−0.406", "+0.433"],
          ["1-D CNN", "0.343", "0.510", "1.031", "−0.434", "+0.121"]]
    tabla(slide, MX, 3.95, 12.1, t4, [3.1, 1.8, 1.8, 1.8, 1.8, 1.8], size=13, resaltar_fila=1)


def b15_busqueda(prs, slide) -> None:
    titulo(slide, "Backup · Hyper-parameter search (Section 8)")
    filas = [["Model", "Trials", "Hyper-param.", "Main search space", "Selected"],
             ["LR", "150", "5", "penalty L1 / L2 / elastic-net, C ∈ [1e-4, 1e3], scaler", "L2, C = 0.000165, robust scaler"],
             ["XGBoost", "150", "9", "trees, depth, learning rate, subsampling, L1/L2", "485 trees, depth 9, lr 0.074"],
             ["LSTM", "80 (54 done)", "17", "layers, units, bidirectional, pooling, look-back, loss", "2 × 118, attention, 10 days, focal loss"],
             ["1-D CNN", "80 (53 done)", "20", "blocks, filters, kernel, dilation, pooling, look-back", "50 → 100 filters, kernel 7, 10 days"],
             ["CNN-LSTM", "80 (74 done)", "19", "conv filters, kernel, LSTM units, look-back", "77 → 154 filters, LSTM 136, 20 days"]]
    tabla(slide, MX, 1.35, 12.1, filas, [1.6, 1.5, 1.4, 4.3, 3.3], size=13)
    texto(slide, MX, 4.9, ANCHO - 2 * MX, 1.2,
          [[("Same TPE sampler (seed 42), same decision year (2024), objective macro-F1. ", {"bold": True}),
            ("One network training takes 7–12 s on a consumer GPU: compute was not the limit. "
             "Symmetric in procedure, not in dimension.", {"color": TINTA2})]],
          size=15)


# ---------------------------------------------------------------------------
def construir() -> Path:
    prs = abrir_plantilla()
    quitar_diapositivas(prs, {0})
    blank = diseno(prs, "Blank")
    cierre = diseno(prs, "3_Diapositiva de título")

    portada = prs.slides[0]
    s01_portada(prs, portada)
    constructores = [s02_pregunta, s03_datos, s04_protocolo, s05_resultado1, s06_resultado2,
                     s07_resultado3, s08_protocolo_corregido, s09_resultado4, s10_por_que]
    slides = [portada]
    for fn in constructores:
        sl = prs.slides.add_slide(blank)
        fn(prs, sl)
        slides.append(sl)
    sl = prs.slides.add_slide(cierre)
    s11_cierre(prs, sl)
    slides.append(sl)
    for fn in (b12_tabla6, b13_despues, b14_tablas, b15_busqueda):
        sl = prs.slides.add_slide(blank)
        fn(prs, sl)
        slides.append(sl)

    for sl, d in zip(slides, DIAPOSITIVAS):
        notas(sl, d)
    for sl, (n, desc) in zip(slides[len(DIAPOSITIVAS):], RESPALDO.items()):
        sl.notes_slide.notes_text_frame.text = f"BACKUP (slide {n}) · show only if a question needs it.\n{desc}"

    prs.core_properties.title = "A Benchmark of Five Machine-Learning Architectures for Daily Trading-Signal Classification on US Tech Equities"
    prs.core_properties.author = "J. B. Polvo Cuatianquiz et al."
    prs.save(str(SALIDA))
    return SALIDA


if __name__ == "__main__":
    print(construir())
